"""Full-fidelity admin Team export; never query private records without authorization.

A spreadsheet consists of one row per registration, a row per team member,
all recorded form answers, and (for Super Admins) payment evidence.  Drafts
are kept distinct from submitted registrations; blanks are never imputed.
"""
from __future__ import annotations

import json
import re
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.repository.query_utils import many

MAX_DETAIL_TEAMS = 1000

# These queries intentionally use a bounded list of registration UUIDs selected by
# the filtered, permission-checked parent query; no free-form SQL is accepted.
MEMBERS_SQL = '''SELECT m.registration_id,m.id AS member_id,m.member_order,m.role,
    m.full_name,m.email,m.mobile,m.photo_file_id,photo.original_filename AS photo_filename,
    r.registration_code,r.team_id,r.team_name,r.status,r.created_at,r.submitted_at,
    e.edition_code,d.domain_name
    FROM public.team_members m
    JOIN public.team_registrations r ON r.id=m.registration_id
    JOIN public.hackathon_editions e ON e.id=r.edition_id
    LEFT JOIN public.domains d ON d.id=r.domain_id
    LEFT JOIN public.file_assets photo ON photo.id=m.photo_file_id
    WHERE m.registration_id = ANY(CAST(:ids AS uuid[]))
    ORDER BY r.created_at DESC,m.member_order'''

ANSWERS_SQL = '''SELECT a.registration_id,a.team_member_id,m.member_order,m.full_name AS member_name,
    m.role AS member_role,q.field_key,q.question_label,q.answer_scope,q.input_type,
    a.answer_json,a.uploaded_file_id,upload_file.original_filename AS answer_filename,
    r.registration_code,r.team_id,r.team_name,r.status,e.edition_code,d.domain_name
    FROM public.form_answers a
    JOIN public.team_registrations r ON r.id=a.registration_id
    JOIN public.form_questions q ON q.id=a.question_id
    JOIN public.hackathon_editions e ON e.id=r.edition_id
    LEFT JOIN public.domains d ON d.id=r.domain_id
    LEFT JOIN public.team_members m ON m.id=a.team_member_id
    LEFT JOIN public.file_assets upload_file ON upload_file.id=a.uploaded_file_id
    WHERE a.registration_id = ANY(CAST(:ids AS uuid[]))
    ORDER BY r.created_at DESC,q.display_order,m.member_order NULLS FIRST'''

PAYMENTS_SQL = '''SELECT p.registration_id,r.registration_code,r.team_id,r.team_name,
    p.transaction_id,p.claimed_amount,p.currency,p.status AS payment_status,
    p.proof_file_id,proof.original_filename AS proof_filename,
    p.submitted_at AS evidence_submitted_at,
    p.verified_at,p.rejection_reason,e.edition_code
    FROM public.payment_submissions p
    JOIN public.team_registrations r ON r.id=p.registration_id
    JOIN public.hackathon_editions e ON e.id=r.edition_id
    LEFT JOIN public.file_assets proof ON proof.id=p.proof_file_id
    WHERE p.registration_id = ANY(CAST(:ids AS uuid[]))
    ORDER BY p.submitted_at DESC NULLS LAST'''

CONSENTS_SQL = '''SELECT m.registration_id,r.registration_code,r.team_id,r.team_name,
    m.member_order,m.full_name AS member_name,c.consent_type,c.response,
    c.policy_version,c.recorded_at,q.question_label
    FROM public.participant_consents c
    JOIN public.team_members m ON m.id=c.member_id
    JOIN public.team_registrations r ON r.id=m.registration_id
    LEFT JOIN public.form_questions q ON q.id=c.question_id
    WHERE m.registration_id=ANY(CAST(:ids AS uuid[]))
    ORDER BY r.created_at DESC,m.member_order,c.recorded_at DESC'''
CONSENT_COLUMNS = ['team_id','registration_code','team_name','member_order',
    'member_name','consent_type','response','policy_version','recorded_at','question_label']

TEAM_COLUMNS = [
    'team_id','registration_code','team_name','edition_code',
    'domain_name','status','members','investor_pitch_interest',
    'created_at','submitted_at',
    'leader_name','leader_email','leader_mobile',
    'all_member_names','all_member_emails','all_member_phones','member_details_json',
]
MEMBER_COLUMNS = [
    'team_id','registration_code','team_name','edition_code','domain_name','status',
    'member_order','role','full_name','email','mobile','photo_filename','photo_file_id',
    'member_id','created_at','submitted_at',
]
ANSWER_COLUMNS = [
    'team_id','registration_code','team_name','edition_code','domain_name','status',
    'member_order','member_name','member_role','answer_scope','question_label',
    'field_key','input_type','answer','answer_json','answer_filename','uploaded_file_id',
]
PAYMENT_COLUMNS = [
    'team_id','registration_code','team_name','edition_code','transaction_id',
    'claimed_amount','currency','payment_status','proof_filename','proof_file_id',
    'evidence_submitted_at','verified_at','rejection_reason',
]


def readable_answer(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except (ValueError, TypeError):
            return value
    if isinstance(value, bool):
        return 'Yes' if value else 'No'
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def raw_answer(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, default=str)


def answer_export_key(answer: dict) -> str:
    """Stable, distinct columns for team answers and each member's answers."""
    field = re.sub(r'[^a-z0-9_]+', '_', str(answer.get('field_key') or '').lower())[:88]
    order = answer.get('member_order')
    prefix = f'member_{order}__' if order is not None else 'team__'
    return 'answer__' + prefix + field


def _id(record: dict) -> str:
    return str(record['registration_id'])


async def fetch_team_details(db, teams: list[dict], *, payments: bool):
    """Fetch every member/answer for exactly the registration rows being exported."""
    ids = [t['registration_id'] for t in teams if t.get('registration_id')]
    if not ids:
        return [], [], [], []
    if len(ids) > MAX_DETAIL_TEAMS:
        from fastapi import HTTPException
        raise HTTPException(422,
            f'A full-detail export supports {MAX_DETAIL_TEAMS} teams per file. '
            'Filter by event, status or date to download the remaining teams.')
    params = {'ids': ids}
    members = await many(db, MEMBERS_SQL, params)
    answers = await many(db, ANSWERS_SQL, params)
    consents = await many(db, CONSENTS_SQL, params)
    evidence = await many(db, PAYMENTS_SQL, params) if payments else []
    return members, answers, consents, evidence


def assemble_teams(teams: list[dict], members: list[dict]) -> list[dict]:
    by_registration: dict[str, list[dict]] = {}
    for member in members:
        by_registration.setdefault(_id(member), []).append(member)
    result = []
    for team in teams:
        row = dict(team)
        people = by_registration.get(_id(team), []) if team.get('registration_id') else []
        people = sorted(people, key=lambda p: int(p.get('member_order') or 999))
        row['all_member_names'] = '\n'.join(str(p.get('full_name') or '') for p in people)
        row['all_member_emails'] = '\n'.join(str(p.get('email') or '') for p in people)
        row['all_member_phones'] = '\n'.join(str(p.get('mobile') or '') for p in people)
        row['member_details_json'] = json.dumps([
            {k: str(p[k]) if p.get(k) is not None else None
             for k in ('member_order','role','full_name','email','mobile','photo_filename','photo_file_id')}
            for p in people], ensure_ascii=False)
        for person in people:
            pos = person.get('member_order')
            if not isinstance(pos, int) or not 1 <= pos <= 20:
                continue
            for field in ('full_name','email','mobile','photo_filename','photo_file_id'):
                row[f'member_{pos}_{field}'] = person.get(field)
            if pos == 1:
                row.update(leader_name=person.get('full_name'),
                           leader_email=person.get('email'),
                           leader_mobile=person.get('mobile'))
        result.append(row)
    return result


def assemble_answers(answers: list[dict]) -> list[dict]:
    result = []
    for answer in answers:
        row = dict(answer)
        row['answer'] = readable_answer(answer.get('answer_json'))
        row['answer_json'] = raw_answer(answer.get('answer_json'))
        result.append(row)
    return result


def csv_flatten(teams: list[dict], answers: list[dict],
                consents: list[dict], base_columns: list[str]):
    """One row per team, with member columns and every recorded answer as columns."""
    unique_answer_keys = sorted({answer_export_key(a) for a in answers})
    consent_keys = sorted({
        f"consent__member_{c.get('member_order')}__{str(c.get('consent_type') or '').lower()}"
        for c in consents if c.get('member_order') is not None
    })
    # Selected base columns are respected; detail fields are intentionally
    # included because this is the full Teams export the user selected.
    member_fields = sorted({k for t in teams for k in t
        if re.fullmatch(r'member_\d+_(full_name|email|mobile|photo_filename|photo_file_id)', k)})
    extended = [*base_columns,
        *(k for k in ('leader_name','leader_email','leader_mobile')
          if k not in base_columns),
        *(k for k in member_fields if k not in base_columns),
        *unique_answer_keys,
        *consent_keys]
    by_registration: dict[str, list[dict]] = {}
    for answer in answers:
        by_registration.setdefault(_id(answer), []).append(answer)
    consents_by_reg: dict[str,list[dict]] = {}
    for consent in consents:
        consents_by_reg.setdefault(_id(consent),[]).append(consent)
    rows = []
    for team in teams:
        row = dict(team)
        for a in by_registration.get(_id(team), []) if team.get('registration_id') else []:
            key = answer_export_key(a)
            val = a['answer'] if a['answer'] else str(a.get('answer_filename') or a.get('uploaded_file_id') or '')
            # Duplicate keys within the same team get combined, not silently lost.
            if key in row and row[key] not in (None, ''):
                if val and val != row[key]:
                    row[key] = str(row[key]) + ' | ' + val
            else:
                row[key] = val
        for consent in consents_by_reg.get(_id(team),[]) if team.get('registration_id') else []:
            key=f"consent__member_{consent.get('member_order')}__{str(consent.get('consent_type') or '').lower()}"
            row[key]=readable_answer(consent.get('response'))
        rows.append(row)
    return extended, rows


def _write_sheet(wb: Workbook, name: str, columns: list[str], rows: list[dict], safe_cell):
    ws = wb.create_sheet(title=name)
    ws.append([c.replace('_',' ').title() for c in columns])
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 34
    fill = PatternFill(fill_type='solid',fgColor='7845D5')
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(bold=True,color='FFFFFF')
        cell.alignment = Alignment(horizontal='center',vertical='center',wrap_text=True)
    for record in rows:
        ws.append([safe_cell(record.get(c)) for c in columns])
    for col, key in enumerate(columns,1):
        sample = [str(safe_cell(r.get(key))) for r in rows[:250]]
        longest_line = max([len(key)] + [max((len(line) for line in value.split('\n')),default=0) for value in sample])
        preferred = 42 if key in ('all_member_names','all_member_emails','all_member_phones','member_details_json') else 34
        ws.column_dimensions[get_column_letter(col)].width = min(preferred,max(16,longest_line+3))
    for row_no in range(2,ws.max_row+1):
        estimated_lines=1
        for col_no, key in enumerate(columns,1):
            cell=ws.cell(row_no,col_no)
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            text=str(cell.value or '')
            width=ws.column_dimensions[get_column_letter(col_no)].width or 22
            import math
            wraps=sum(max(1,math.ceil(len(line)/max(1,width-3))) for line in text.split('\n'))
            estimated_lines=max(estimated_lines,wraps)
        ws.row_dimensions[row_no].height=min(240,max(24,estimated_lines*16+8))
    ws.auto_filter.ref=f'A1:{get_column_letter(len(columns))}{max(1,ws.max_row)}'
    return ws


def render_team_workbook(teams: list[dict], members: list[dict], answers: list[dict],
                         consents: list[dict], evidence: list[dict], columns: list[str], safe_cell,
                         *, super_admin: bool) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    team_columns = list(dict.fromkeys(columns))
    _write_sheet(wb,'Teams',team_columns,teams,safe_cell)
    _write_sheet(wb,'Members',MEMBER_COLUMNS,members,safe_cell)
    _write_sheet(wb,'Registration Answers',ANSWER_COLUMNS,answers,safe_cell)
    _write_sheet(wb,'Consents',CONSENT_COLUMNS,consents,safe_cell)
    if super_admin:
        _write_sheet(wb,'Payment Evidence',PAYMENT_COLUMNS,evidence,safe_cell)
    buffer=BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def render_team_json(teams: list[dict], members: list[dict], answers: list[dict],
                     consents: list[dict], evidence: list[dict], *, super_admin: bool) -> bytes:
    """Complete structured backup of selected registrations; no invented data."""
    by_id = lambda records: {str(t['registration_id']): [dict(v) for v in records if str(v.get('registration_id')) == str(t['registration_id'])] for t in teams}
    by_member = by_id(members)
    by_answer = by_id(answers)
    by_consent = by_id(consents)
    by_evidence = by_id(evidence) if super_admin else {}
    result = []
    for team in teams:
        rid = str(team['registration_id'])
        record = {'team':dict(team), 'members':by_member[rid], 'registration_answers':by_answer[rid],
                  'consents':by_consent[rid]}
        for answer in record['registration_answers']:
            answer['answer_json'] = _parse_answer(answer.get('answer_json'))
        if super_admin: record['payment_evidence'] = by_evidence.get(rid,[])
        result.append(record)
    return json.dumps({'schema_version':1,'registrations':result},ensure_ascii=False,indent=2,default=str).encode('utf-8')


def _parse_answer(value):
    if isinstance(value,str):
        try:return json.loads(value)
        except (TypeError, ValueError):return value
    return value


def render_members_workbook(members: list[dict], answers: list[dict],
                            consents: list[dict], columns: list[str], safe_cell) -> bytes:
    wb=Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, 'Members', columns, members, safe_cell)
    _write_sheet(wb, 'Registration Answers', ANSWER_COLUMNS, answers, safe_cell)
    _write_sheet(wb, 'Consents', CONSENT_COLUMNS, consents, safe_cell)
    buffer=BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def csv_flatten_members(members: list[dict], answers: list[dict],
                        consents: list[dict], base_columns: list[str]):
    """One row per member with their own answers and applicable team answers."""
    keys=sorted({answer_export_key(a) for a in answers})
    consent_keys = sorted({
        f"consent__{str(c.get('consent_type') or '').lower()}" for c in consents
    })
    cols=[*base_columns, *(k for k in keys if k not in base_columns),*consent_keys]
    consents_by_reg: dict[str,list[dict]]={}
    for c in consents:
        consents_by_reg.setdefault(_id(c),[]).append(c)
    by_registration: dict[str,list[dict]]={}
    for a in answers:
        by_registration.setdefault(_id(a),[]).append(a)
    result=[]
    for member in members:
        row=dict(member)
        for a in by_registration.get(_id(member),[]):
            if a.get('member_order') is not None and a.get('member_order')!=member.get('member_order'):
                continue
            k=answer_export_key(a)
            value=a.get('answer') or str(a.get('answer_filename') or a.get('uploaded_file_id') or '')
            if value:
                row[k]=value
        for c in consents_by_reg.get(_id(member),[]):
            if c.get('member_order')==member.get('member_order'):
                row[f"consent__{str(c.get('consent_type') or '').lower()}"]=readable_answer(c.get('response'))
        result.append(row)
    return cols,result
