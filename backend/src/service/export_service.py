"""Filtered XLSX/CSV exports. Purple headers and Excel AutoFilters; no raw SQL from clients."""
from csv import writer
from io import StringIO,BytesIO
from uuid import UUID, uuid4
from datetime import datetime
from fastapi import HTTPException,UploadFile
from starlette.datastructures import Headers
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from src.repository.query_utils import many,one
from src.service.access_service import require, require_super
from src.service.storage_service import upload, EXPORT,signed_url
from src.service import team_export_details as team_details

DATASETS={
 'TEAMS':('TEAMS_VIEW', '''SELECT r.id AS registration_id,r.team_id,r.registration_code,r.team_name,r.source_team_id,r.status,
    r.participant_category,r.investor_pitch_interest,r.created_at,r.submitted_at,
    e.edition_code,d.domain_name,
    (SELECT count(*) FROM public.team_members m WHERE m.registration_id=r.id) AS members
    FROM public.team_registrations r
    JOIN public.hackathon_editions e ON e.id=r.edition_id
    LEFT JOIN public.domains d ON d.id=r.domain_id
    WHERE (CAST(:edition AS uuid) IS NULL OR r.edition_id=:edition)
    AND ((CAST(:status AS text) IS NULL AND r.status<>'DRAFT') OR r.status=:status)
    AND (CAST(:domain AS uuid) IS NULL OR r.domain_id=:domain)
    AND (CAST(:start AS timestamptz) IS NULL OR r.created_at>=:start)
    AND (CAST(:end AS timestamptz) IS NULL OR r.created_at<=:end)
    ORDER BY r.created_at DESC LIMIT 10001'''),
 'MEMBERS':('PARTICIPANTS_VIEW', '''SELECT m.registration_id,r.team_id,r.registration_code,r.team_name,m.member_order,
    m.id AS member_id,m.role,m.full_name,m.email,m.mobile,m.photo_file_id,photo.original_filename AS photo_filename,
    r.status,e.edition_code,d.domain_name,r.created_at,r.submitted_at FROM public.team_members m
    JOIN public.team_registrations r ON r.id=m.registration_id
    JOIN public.hackathon_editions e ON e.id=r.edition_id
    LEFT JOIN public.domains d ON d.id=r.domain_id
    LEFT JOIN public.file_assets photo ON photo.id=m.photo_file_id
    WHERE (CAST(:edition AS uuid) IS NULL OR r.edition_id=:edition)
    AND ((CAST(:status AS text) IS NULL AND r.status<>'DRAFT') OR r.status=:status)
    AND (CAST(:domain AS uuid) IS NULL OR r.domain_id=:domain)
    AND (CAST(:start AS timestamptz) IS NULL OR r.created_at>=:start)
    AND (CAST(:end AS timestamptz) IS NULL OR r.created_at<=:end)
    ORDER BY r.created_at DESC,m.member_order LIMIT 10001'''),
 'PAYMENTS':('PAYMENT_VIEW','''SELECT r.team_id,r.registration_code,r.team_name,p.transaction_id,
    p.proof_file_id,proof.original_filename AS proof_filename,p.claimed_amount,p.currency,p.status,p.submitted_at,p.verified_at,
    p.rejection_reason,e.edition_code,d.domain_name
    FROM public.payment_submissions p
    JOIN public.team_registrations r ON r.id=p.registration_id
    JOIN public.hackathon_editions e ON e.id=r.edition_id
    LEFT JOIN public.domains d ON d.id=r.domain_id
    LEFT JOIN public.file_assets proof ON proof.id=p.proof_file_id
    WHERE r.status='APPROVED' AND p.status='VERIFIED' AND e.edition_code NOT LIKE 'MIIH_TEST_%'
    AND (CAST(:edition AS uuid) IS NULL OR r.edition_id=:edition)
    AND (CAST(:status AS text) IS NULL OR p.status=:status)
    AND (CAST(:domain AS uuid) IS NULL OR r.domain_id=:domain)
    AND (CAST(:start AS timestamptz) IS NULL OR p.submitted_at>=:start)
    AND (CAST(:end AS timestamptz) IS NULL OR p.submitted_at<=:end)
    ORDER BY p.submitted_at DESC LIMIT 10001'''),
 'STATISTICS':('STATISTICS_VIEW', '''SELECT d.domain_name,e.edition_code,
   COUNT(DISTINCT r.id) FILTER (WHERE r.status<>'DRAFT') AS registered_teams,
   COUNT(DISTINCT m.id) FILTER (WHERE r.status<>'DRAFT') AS registered_participants
   FROM public.team_registrations r
   JOIN public.hackathon_editions e ON e.id=r.edition_id
   JOIN public.domains d ON d.id=r.domain_id
   LEFT JOIN public.team_members m ON m.registration_id=r.id
   WHERE (CAST(:edition AS uuid) IS NULL OR r.edition_id=:edition)
     AND ((CAST(:status AS text) IS NULL AND r.status<>'DRAFT') OR r.status=:status)
     AND (CAST(:domain AS uuid) IS NULL OR r.domain_id=:domain)
     AND (CAST(:start AS timestamptz) IS NULL OR r.created_at>=:start)
     AND (CAST(:end AS timestamptz) IS NULL OR r.created_at<=:end)
   GROUP BY d.domain_name,e.edition_code ORDER BY registered_teams DESC LIMIT 10001'''),
 'EDITION_HISTORY':('HISTORY_VIEW','''SELECT e.edition_code,e.title,e.year,e.status,
    e.starts_at,e.ends_at,e.location,
    COUNT(r.id) FILTER (WHERE r.status<>'DRAFT') AS registered_teams
    FROM public.hackathon_editions e
    LEFT JOIN public.team_registrations r ON r.edition_id=e.id
    WHERE (CAST(:edition AS uuid) IS NULL OR e.id=:edition)
      AND (CAST(:status AS text) IS NULL OR e.status=:status)
      AND (CAST(:start AS timestamptz) IS NULL OR e.starts_at>=:start)
      AND (CAST(:end AS timestamptz) IS NULL OR e.starts_at<=:end)
    GROUP BY e.id ORDER BY e.starts_at DESC NULLS LAST LIMIT 10001'''),
 'ANSWERS':('PARTICIPANTS_VIEW','''SELECT r.team_id,r.registration_code,r.team_name,
    m.member_order,m.full_name AS member_name,m.email AS member_email,m.role AS member_role,q.field_key,q.question_label,q.answer_scope,
    a.answer_json::text AS answer,a.uploaded_file_id::text AS file_asset_id,upload_file.original_filename AS answer_filename,
    e.edition_code,r.status,r.created_at,d.domain_name
    FROM public.form_answers a
    JOIN public.team_registrations r ON r.id=a.registration_id
    JOIN public.form_questions q ON q.id=a.question_id
    JOIN public.hackathon_editions e ON e.id=r.edition_id
    LEFT JOIN public.domains d ON d.id=r.domain_id
    LEFT JOIN public.team_members m ON m.id=a.team_member_id
    LEFT JOIN public.file_assets upload_file ON upload_file.id=a.uploaded_file_id
    WHERE (CAST(:edition AS uuid) IS NULL OR r.edition_id=:edition)
    AND ((CAST(:status AS text) IS NULL AND r.status<>'DRAFT') OR r.status=:status)
    AND (CAST(:domain AS uuid) IS NULL OR r.domain_id=:domain)
    AND (CAST(:start AS timestamptz) IS NULL OR r.created_at>=:start)
    AND (CAST(:end AS timestamptz) IS NULL OR r.created_at<=:end)
    ORDER BY r.created_at DESC,q.display_order LIMIT 10001''')
}
DEFAULT_COLS={
 'TEAMS':team_details.TEAM_COLUMNS,
 'MEMBERS':['team_id','registration_code','team_name','member_order','role','full_name','email','mobile','photo_filename','photo_file_id','domain_name','status','created_at','submitted_at'],
 'PAYMENTS':['team_id','registration_code','team_name','transaction_id','claimed_amount','currency','status','proof_filename','proof_file_id','submitted_at','verified_at','rejection_reason'],
 'STATISTICS':['domain_name','registered_teams','registered_participants','edition_code'],
 'EDITION_HISTORY':['edition_code','title','year','status','starts_at','ends_at','location','registered_teams'],
 'ANSWERS':['team_id','registration_code','team_name','member_order','member_name','member_email','field_key','question_label','answer_scope','answer','answer_filename','file_asset_id']
}

def safe_cell(value):
    if value is None:return ''
    if isinstance(value,(dict,list)):value=str(value)
    if isinstance(value,datetime):return value.isoformat()
    if isinstance(value,UUID):return str(value)
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@','\t','\r')):
        return "'"+value
    return value

def render_excel(columns,records):
    wb=Workbook();ws=wb.active;ws.title='Made In India'
    ws.append([c.replace('_',' ').title() for c in columns]);ws.freeze_panes='A2'
    header=PatternFill(fill_type='solid',fgColor='8C52FF')
    for cell in ws[1]:
        cell.fill=header;cell.font=Font(color='FFFFFF',bold=True)
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    ws.row_dimensions[1].height=32
    for r in records:
        ws.append([safe_cell(r.get(c)) for c in columns])
    for row in ws.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
    for j,c in enumerate(columns,1):
        max_width=max(len(str(c)),*(len(str(safe_cell(x.get(c)))) for x in records[:100])) if records else len(c)
        ws.column_dimensions[get_column_letter(j)].width=min(55,max(18,max_width+3))
    ws.auto_filter.ref=f'A1:{get_column_letter(len(columns))}{max(1,ws.max_row)}'
    b=BytesIO();wb.save(b);return b.getvalue()

def render_csv(columns,records):
    output=StringIO(newline='');w=writer(output)
    w.writerow([c.replace('_',' ').title() for c in columns])
    for r in records:w.writerow([safe_cell(r.get(c)) for c in columns])
    return ('\ufeff'+output.getvalue()).encode('utf-8')

async def run_export(db,admin,payload):
    await require(db,admin,'EXPORT_CREATE')
    permission,sql=DATASETS[payload.dataset]
    if payload.dataset=='PAYMENTS':require_super(admin)
    elif payload.dataset=='TEAMS':await require(db,admin,permission,'PARTICIPANTS_VIEW')
    else:await require(db,admin,permission)
    if payload.file_format=='JSON' and payload.dataset!='TEAMS':
        raise HTTPException(422,'JSON is available for the Teams dataset')
    if payload.date_from and payload.date_to and payload.date_from>payload.date_to:
        raise HTTPException(422,'Invalid date range')
    params={'edition':payload.edition_id,'domain':payload.domain_id,
            'status':payload.status,'start':payload.date_from,'end':payload.date_to}
    records=await many(db,sql,params)
    if len(records)>10000:raise HTTPException(422,'Export exceeds 10,000 rows. Narrow the filters.')
    all_columns=set(records[0]) if records else set(DEFAULT_COLS[payload.dataset])
    if payload.dataset=='TEAMS':all_columns.update(team_details.TEAM_COLUMNS)
    columns=payload.columns or DEFAULT_COLS[payload.dataset]
    if len(set(columns))!=len(columns) or any(col not in all_columns for col in columns):
        raise HTTPException(422,'Requested column is not available')
    if not columns:raise HTTPException(422,'Select one or more columns')
    if payload.dataset == 'TEAMS':
        members, answers, consents, evidence = await team_details.fetch_team_details(
            db, records, payments=admin['role']=='SUPER_ADMIN')
        expanded = team_details.assemble_teams(records,members)
        answers = team_details.assemble_answers(answers)
        if payload.file_format=='JSON':
            raw=team_details.render_team_json(expanded,members,answers,consents,evidence,super_admin=admin['role']=='SUPER_ADMIN')
        elif payload.file_format=='XLSX':
            raw=team_details.render_team_workbook(expanded,members,answers,consents,evidence,columns,
                safe_cell,super_admin=admin['role']=='SUPER_ADMIN')
        else:
            flat_cols,flat_rows=team_details.csv_flatten(expanded,answers,consents,columns)
            raw=render_csv(flat_cols,flat_rows)
    elif payload.dataset=='MEMBERS':
        identifiers=list(dict.fromkeys(r['registration_id'] for r in records if r.get('registration_id')))
        _, raw_answers, consents, _ = await team_details.fetch_team_details(
            db,[{'registration_id':i} for i in identifiers],payments=False)
        answer_rows=team_details.assemble_answers(raw_answers)
        if payload.file_format=='XLSX':
            raw=team_details.render_members_workbook(records,answer_rows,consents,columns,safe_cell)
        else:
            flat_cols,flat_rows=team_details.csv_flatten_members(records,answer_rows,consents,columns)
            raw=render_csv(flat_cols,flat_rows)
    else:
        raw=render_excel(columns,records) if payload.file_format=='XLSX' else render_csv(columns,records)
    mime='application/json' if payload.file_format=='JSON' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if payload.file_format=='XLSX' else 'text/csv'
    filename=f'{payload.dataset.lower()}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.{payload.file_format.lower()}'
    file=UploadFile(file=BytesIO(raw),filename=filename,headers=Headers({'content-type':mime}))
    asset=await upload(db,bucket=EXPORT,file=file,edition_id=payload.edition_id,admin_id=admin['id'])
    import json
    record=await one(db,'''INSERT INTO public.export_runs
      (requested_by_admin_id,edition_id,dataset,file_format,applied_filters,
       selected_columns,row_count,output_file_id,status,completed_at,created_by)
      VALUES (:actor,:edition,:dataset,:format,cast(:filters AS jsonb),cast(:cols AS jsonb),
        :rows,:file,'DONE',now(),:actor)
      RETURNING id,dataset,file_format,row_count,status''',
      {'actor':admin['id'],'edition':payload.edition_id,'dataset':payload.dataset,
       'format':payload.file_format,'filters':json.dumps(payload.model_dump(mode='json',exclude={'columns'})),
       'cols':json.dumps(columns),'rows':len(records),'file':asset['id']})
    await db.commit()
    record['download_url']=await signed_url(asset)
    record['expires_in_seconds']=120
    return record

async def history(db,admin):
    await require(db,admin,'EXPORT_VIEW')
    where='' if admin['role']=='SUPER_ADMIN' else 'WHERE r.requested_by_admin_id=:actor'
    return await many(db,'''SELECT r.id,r.dataset,r.file_format,r.row_count,r.status,
      r.requested_at,r.completed_at,r.selected_columns,r.applied_filters
      FROM public.export_runs r '''+where+' ORDER BY requested_at DESC LIMIT 100',{'actor':admin['id']})


async def download_direct(db, admin, payload):
    # Direct exports require admin access; private image uploads still use Storage.
    await require(db, admin, 'EXPORT_CREATE')
    permission, sql = DATASETS[payload.dataset]
    if payload.dataset == 'PAYMENTS':
        require_super(admin)
    elif payload.dataset == 'TEAMS':
        await require(db, admin, permission, 'PARTICIPANTS_VIEW')
    else:
        await require(db, admin, permission)
    if payload.file_format=='JSON' and payload.dataset!='TEAMS':
        raise HTTPException(422,'JSON is available for the Teams dataset')
    if payload.date_from and payload.date_to and payload.date_from > payload.date_to:
        raise HTTPException(422, 'Invalid date range')
    params = {'edition':payload.edition_id,'domain':payload.domain_id,
              'status':payload.status,'start':payload.date_from,'end':payload.date_to}
    records = await many(db, sql, params)
    if len(records) > 10000:
        raise HTTPException(422, 'Export exceeds 10,000 rows. Narrow your filters.')
    allowed = set(records[0]) if records else set(DEFAULT_COLS[payload.dataset])
    if payload.dataset=='TEAMS':allowed.update(team_details.TEAM_COLUMNS)
    columns = payload.columns or DEFAULT_COLS[payload.dataset]
    if not columns or len(set(columns)) != len(columns) or any(c not in allowed for c in columns):
        raise HTTPException(422, 'Requested columns are not available')
    if payload.dataset=='TEAMS':
        members, answers, consents, evidence = await team_details.fetch_team_details(
            db,records,payments=admin['role']=='SUPER_ADMIN')
        expanded=team_details.assemble_teams(records,members)
        answers=team_details.assemble_answers(answers)
        if payload.file_format=='JSON':
            raw=team_details.render_team_json(expanded,members,answers,consents,evidence,super_admin=admin['role']=='SUPER_ADMIN')
        elif payload.file_format=='XLSX':
            raw=team_details.render_team_workbook(expanded,members,answers,consents,evidence,
                columns,safe_cell,super_admin=admin['role']=='SUPER_ADMIN')
        else:
            flat_cols,flat_rows=team_details.csv_flatten(expanded,answers,consents,columns)
            raw=render_csv(flat_cols,flat_rows)
    elif payload.dataset=='MEMBERS':
        identifiers=list(dict.fromkeys(r['registration_id'] for r in records if r.get('registration_id')))
        _,raw_answers,consents,_=await team_details.fetch_team_details(
            db,[{'registration_id':i} for i in identifiers],payments=False)
        answer_rows=team_details.assemble_answers(raw_answers)
        if payload.file_format=='XLSX':
            raw=team_details.render_members_workbook(records,answer_rows,consents,columns,safe_cell)
        else:
            flat_cols,flat_rows=team_details.csv_flatten_members(records,answer_rows,consents,columns)
            raw=render_csv(flat_cols,flat_rows)
    else:
        raw = render_excel(columns,records) if payload.file_format=='XLSX' else render_csv(columns,records)
    mime = ('application/json; charset=utf-8' if payload.file_format=='JSON' else
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if payload.file_format=='XLSX' else 'text/csv; charset=utf-8')
    filename = f'miih_{payload.dataset.lower()}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.{payload.file_format.lower()}'
    import json
    await one(db, '''INSERT INTO public.export_runs
        (requested_by_admin_id,edition_id,dataset,file_format,applied_filters,
         selected_columns,row_count,output_file_id,status,completed_at,created_by)
        VALUES (:actor,:edition,:dataset,:format,cast(:filters AS jsonb),
                cast(:cols AS jsonb),:rows,NULL,'DONE',now(),:actor)
        RETURNING id''', {
          'actor':admin['id'],'edition':payload.edition_id,'dataset':payload.dataset,
          'format':payload.file_format,'filters':json.dumps(payload.model_dump(mode='json',exclude={'columns'})),
          'cols':json.dumps(columns),'rows':len(records)})
    await db.commit()
    return raw, mime, filename, len(records)
