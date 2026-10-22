"""Offline checks; deliberately never contact user's Supabase database."""
from uuid import uuid4
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from main import app
from src.service.settings_service import safe_payment_link,DUMMY_URL
from src.service.export_service import render_excel,render_csv,DEFAULT_COLS
from src.service.registration_service import verify_draft_token
from src.service.access_service import require_super
from openpyxl import load_workbook
from io import BytesIO
import hashlib


def test_dummy_url_cannot_be_activated():
    assert not safe_payment_link(DUMMY_URL,True)
    assert safe_payment_link(DUMMY_URL,False)
    assert not safe_payment_link(None,True)
    assert safe_payment_link('https://rzp.io/i/abc123',True)
    assert safe_payment_link('https://pages.razorpay.com/madeinindia',True)

@pytest.mark.parametrize('url',["https://evil.example/miih","http://rzp.io/a","https://rzp.io.evil.example/x",
                               "javascript:alert(1)","https://example.com/miih-payment-placeholder"])
def test_invalid_active_payment_hosts_rejected(url):
    assert not safe_payment_link(url,True)

def test_admin_cannot_become_super_admin():
    with pytest.raises(HTTPException) as e:require_super({'role':'ADMIN'})
    assert e.value.status_code==403

def test_registration_edit_token_is_hashed_and_required():
    token='A'*45
    record={'id':uuid4(),'leader_edit_token_hash':hashlib.sha256(token.encode()).hexdigest()}
    verify_draft_token(record,token)
    with pytest.raises(HTTPException):verify_draft_token(record,'B'*45)
    with pytest.raises(HTTPException):verify_draft_token(record,None)

def test_xlsx_header_is_purple_and_filterable():
    b=render_excel(['team_name','status'],[{'team_name':'Test team','status':'SUBMITTED'}])
    ws=load_workbook(BytesIO(b)).active
    assert ws['A1'].fill.fgColor.rgb.endswith('8C52FF')
    assert ws['A1'].font.bold
    assert ws.freeze_panes=='A2'
    assert ws.auto_filter.ref=='A1:B2'

def test_export_prevents_formula_injection():
    b=render_csv(['team_name'],[{'team_name':'=IMPORTXML("https://evil.test")'}])
    assert b"'=IMPORTXML" in b

def test_public_paths_and_admin_paths_exposed():
    routes=set(app.openapi()['paths'])
    for p in ('/health','/api/v1/auth/login','/api/v1/auth/signup',
              '/api/v1/admin/apply','/api/v1/admin/applications/pending',
              '/api/v1/admin/editions/{edition_id}/payments/config',
              '/api/v1/admin/cms/preview/{page_key}',
              '/api/v1/admin/exports',
              '/api/v1/public/editions/{code}/registration-config',
              '/api/v1/registrations/{registration_id}/submit'):
        assert p in routes,p

def test_restricted_admin_routes_need_authentication():
    c=TestClient(app)
    assert c.get('/api/v1/admin/dashboard').status_code==401
    assert c.post('/api/v1/admin/exports',json={'dataset':'TEAMS'}).status_code==401
    assert c.get('/api/v1/admin/cms/preview/HOME').status_code==401
    assert c.post('/api/v1/admin/editions/'+str(uuid4())+'/payments/config',json={}).status_code==401
