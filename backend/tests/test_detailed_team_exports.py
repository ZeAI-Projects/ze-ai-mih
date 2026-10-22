"""Regression tests for full-data Teams exports (including unfinished drafts)."""
from io import BytesIO
from uuid import UUID
from unittest.mock import AsyncMock
import csv
from io import StringIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from src.models.admin_ops_requests import ExportRequest
from src.service import export_service as service
from src.service import team_export_details as detail

RID=UUID('647a803f-f51b-4616-91c6-043b9848aef0')
TEAM={
    'registration_id':RID,'team_id':None,'registration_code':'MIIH-2F89748BB59247',
    'team_name':'ZeAI','edition_code':'MIIH_UPCOMING','domain_name':'AI and automation',
    'status':'DRAFT','members':4,'investor_pitch_interest':False,'submitted_at':None,
    'created_at':None,
}
PEOPLE=[{'registration_id':RID,'registration_code':TEAM['registration_code'],
         'member_order':i,'role':'LEADER' if i==1 else 'MEMBER',
         'full_name':f'Tester {i}','email':f'tester{i}@example.com',
         'mobile':f'12345000{i:02d}','photo_file_id':UUID(int=i+100),'photo_filename':f'photo{i}.jpg',
         'team_name':'ZeAI','team_id':None,'edition_code':'MIIH_UPCOMING',
         'domain_name':'AI and automation','status':'DRAFT','created_at':None,
         'submitted_at':None,'member_id':f'm-{i}'} for i in range(1,5)]
ANSWERS=[
    {'registration_id':RID,'team_member_id':None,'member_order':None,
     'field_key':'college_name','question_label':'Which college do you study at?',
     'answer_scope':'TEAM','input_type':'TEXT','answer_json':'"Trichy Engineering"',
     'uploaded_file_id':None,'registration_code':TEAM['registration_code']},
    {'registration_id':RID,'team_member_id':None,'member_order':None,
     'field_key':'investor_pitch_interest','question_label':'Interested in investor pitch?',
     'answer_scope':'TEAM','input_type':'BOOLEAN','answer_json':'false',
     'uploaded_file_id':None,'registration_code':TEAM['registration_code']},
]
PAYMENTS=[{'registration_id':RID,'registration_code':TEAM['registration_code'],
           'transaction_id':'PAY-0001','payment_status':'PENDING','claimed_amount':3000,
           'proof_file_id':UUID(int=350)}]

class DummyDb:
    def __init__(self):self.committed=False
    async def commit(self):self.committed=True


def patch_db(monkeypatch):
    calls=[]
    async def dataset_query(db,sql,params):
        calls.append(sql)
        if 'FROM public.team_members m\n    JOIN public.team_registrations' in sql:return PEOPLE
        if 'FROM public.form_answers a' in sql:return ANSWERS
        if 'FROM public.payment_submissions p' in sql:return PAYMENTS
        if 'FROM public.participant_consents c' in sql:return []
        return [TEAM]
    monkeypatch.setattr(service,'many',dataset_query)
    monkeypatch.setattr(detail,'many',dataset_query)
    monkeypatch.setattr(service,'one',AsyncMock(return_value={'id':'export-history'}))
    return calls

@pytest.mark.asyncio
async def test_team_xlsx_has_all_member_details_answers_and_superadmin_payments(monkeypatch):
    patch_db(monkeypatch)
    admin={'role':'SUPER_ADMIN','id':RID}
    raw,mime,name,count=await service.download_direct(DummyDb(),admin,
        ExportRequest(dataset='TEAMS',file_format='XLSX'))
    wb=load_workbook(BytesIO(raw))
    assert wb.sheetnames==['Teams','Members','Registration Answers','Consents','Payment Evidence']
    summary=wb['Teams']
    headers=[cell.value for cell in summary[1]]
    assert summary.cell(2,headers.index('Leader Email')+1).value=='tester1@example.com'
    assert summary.cell(2,headers.index('All Member Names')+1).value=='Tester 1\nTester 2\nTester 3\nTester 4'
    assert summary.cell(2,headers.index('Team Id')+1).value is None  # Not submitted yet
    assert summary.cell(2,headers.index('Submitted At')+1).value is None
    assert wb['Members'].max_row==5
    assert wb['Members']['J2'].value=='tester1@example.com'
    values=[row for row in wb['Registration Answers'].iter_rows(values_only=True)]
    assert any('Which college do you study at?' in row for row in values)
    assert any('Trichy Engineering' in row for row in values)
    assert any('No' in row for row in values)
    assert wb['Payment Evidence'].max_row==2
    assert count==1
    assert name.endswith('.xlsx')

@pytest.mark.asyncio
async def test_team_csv_contains_all_member_email_columns_and_answer_columns(monkeypatch):
    patch_db(monkeypatch)
    raw,_,_,count=await service.download_direct(DummyDb(),{'role':'SUPER_ADMIN','id':RID},
        ExportRequest(dataset='TEAMS',file_format='CSV'))
    rows=list(csv.DictReader(StringIO(raw.decode('utf-8-sig'))))
    assert len(rows)==1
    assert rows[0]['Member 4 Email']=='tester4@example.com'
    assert rows[0]['Answer  Team  College Name']=='Trichy Engineering'
    assert rows[0]['Answer  Team  Investor Pitch Interest']=='No'
    assert count==1

@pytest.mark.asyncio
async def test_team_detail_export_requires_participant_permission(monkeypatch):
    async def require_rights(db,admin,*rights):
        if 'PARTICIPANTS_VIEW' in rights:
            raise HTTPException(403,'Missing permission: PARTICIPANTS_VIEW')
    monkeypatch.setattr(service,'require',require_rights)
    monkeypatch.setattr(service,'many',AsyncMock(side_effect=AssertionError('No query allowed')))
    with pytest.raises(HTTPException) as exc:
        await service.download_direct(DummyDb(),{'role':'ADMIN','id':RID},
            ExportRequest(dataset='TEAMS'))
    assert exc.value.status_code==403

@pytest.mark.asyncio
async def test_regular_admin_does_not_receive_payment_evidence(monkeypatch):
    calls=patch_db(monkeypatch)
    monkeypatch.setattr(service,'require',AsyncMock(return_value=None))
    data,*_=await service.download_direct(DummyDb(),{'role':'ADMIN','id':RID},
        ExportRequest(dataset='TEAMS',file_format='XLSX'))
    wb=load_workbook(BytesIO(data))
    assert 'Members' in wb.sheetnames
    assert 'Payment Evidence' not in wb.sheetnames
    assert not any('FROM public.payment_submissions p' in sql for sql in calls)

@pytest.mark.asyncio
async def test_export_without_matching_registrations_has_headers(monkeypatch):
    monkeypatch.setattr(service,'many',AsyncMock(return_value=[]))
    monkeypatch.setattr(service,'one',AsyncMock(return_value={'id':'audit'}))
    raw,*_=await service.download_direct(DummyDb(),{'role':'SUPER_ADMIN','id':RID},
        ExportRequest(dataset='TEAMS',file_format='XLSX'))
    wb=load_workbook(BytesIO(raw))
    assert wb['Teams'].max_row==1
    assert wb['Members'].max_row==1
    assert wb['Registration Answers'].max_row==1

@pytest.mark.asyncio
async def test_many_detailed_teams_are_limited_before_detail_queries(monkeypatch):
    rows=[{**TEAM,'registration_id':UUID(int=i+1)} for i in range(detail.MAX_DETAIL_TEAMS+1)]
    monkeypatch.setattr(service,'many',AsyncMock(return_value=rows))
    monkeypatch.setattr(detail,'many',AsyncMock(side_effect=AssertionError('Too many teams')))
    with pytest.raises(HTTPException) as exc:
        await service.download_direct(DummyDb(),{'role':'SUPER_ADMIN','id':RID},
            ExportRequest(dataset='TEAMS',file_format='XLSX'))
    assert exc.value.status_code==422

@pytest.mark.asyncio
async def test_members_xlsx_includes_question_answers_in_second_sheet(monkeypatch):
    async def fake_many(db, sql, params):
        if 'FROM public.team_members m\n    JOIN public.team_registrations' in sql:
            return PEOPLE
        if 'FROM public.form_answers a' in sql:
            return ANSWERS
        if 'FROM public.participant_consents c' in sql:
            return []
        return PEOPLE
    monkeypatch.setattr(service,'many',fake_many)
    monkeypatch.setattr(detail,'many',fake_many)
    monkeypatch.setattr(service,'one',AsyncMock(return_value={'id':'audit'}))
    raw,*_=await service.download_direct(DummyDb(),{'role':'SUPER_ADMIN','id':RID},
        ExportRequest(dataset='MEMBERS',file_format='XLSX'))
    wb=load_workbook(BytesIO(raw))
    assert wb.sheetnames == ['Members','Registration Answers','Consents']
    assert wb['Members'].max_row==5
    assert any('Trichy Engineering' in row for row in
               wb['Registration Answers'].iter_rows(values_only=True))

@pytest.mark.asyncio
async def test_members_csv_repeats_team_answers_but_not_other_members_answers(monkeypatch):
    answers=[*ANSWERS,{
        'registration_id':RID, 'member_order':2, 'field_key':'department',
        'question_label':'Department','answer_json':'"CSE"',
        'uploaded_file_id':None}]
    async def fake_many(db, sql, params):
        if 'FROM public.form_answers a' in sql:return answers
        if 'FROM public.participant_consents c' in sql:return []
        return PEOPLE
    monkeypatch.setattr(service,'many',fake_many)
    monkeypatch.setattr(detail,'many',fake_many)
    monkeypatch.setattr(service,'one',AsyncMock(return_value={'id':'audit'}))
    raw,*_=await service.download_direct(DummyDb(),{'role':'SUPER_ADMIN','id':RID},
        ExportRequest(dataset='MEMBERS',file_format='CSV'))
    records=list(csv.DictReader(StringIO(raw.decode('utf-8-sig'))))
    assert len(records)==4
    assert all(r['Answer  Team  College Name']=='Trichy Engineering' for r in records)
    assert records[0]['Answer  Member 2  Department']==''
    assert records[1]['Answer  Member 2  Department']=='CSE'

@pytest.mark.asyncio
async def test_consents_and_original_photo_names_are_exported_when_recorded(monkeypatch):
    consent={'registration_id':RID,'member_order':1,'registration_code':TEAM['registration_code'],
             'consent_type':'TERMS','response':True,'policy_version':'v2',
             'member_name':'Tester 1'}
    async def fake_many(db,sql,params):
        if 'FROM public.team_members m\n    JOIN public.team_registrations' in sql:return PEOPLE
        if 'FROM public.form_answers a' in sql:return ANSWERS
        if 'FROM public.participant_consents c' in sql:return [consent]
        if 'FROM public.payment_submissions p' in sql:return PAYMENTS
        return [TEAM]
    monkeypatch.setattr(service,'many',fake_many)
    monkeypatch.setattr(detail,'many',fake_many)
    monkeypatch.setattr(service,'one',AsyncMock(return_value={'id':'audit'}))
    raw,*_=await service.download_direct(DummyDb(),{'role':'SUPER_ADMIN','id':RID},
        ExportRequest(dataset='TEAMS',file_format='XLSX'))
    wb=load_workbook(BytesIO(raw))
    assert 'Consents' in wb.sheetnames
    assert 'TERMS' in next(wb['Consents'].iter_rows(min_row=2,values_only=True))
    colnames=[c.value for c in wb['Members'][1]]
    assert wb['Members'].cell(2,colnames.index('Photo Filename')+1).value=='photo1.jpg'
    csv_raw,*_=await service.download_direct(DummyDb(),{'role':'SUPER_ADMIN','id':RID},
        ExportRequest(dataset='TEAMS',file_format='CSV'))
    records=list(csv.DictReader(StringIO(csv_raw.decode('utf-8-sig'))))
    assert records[0]['Consent  Member 1  Terms']=='Yes'
