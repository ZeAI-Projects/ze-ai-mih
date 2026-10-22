"""Regression tests for all-in-one staff management and grouped team exports."""
import json
from io import BytesIO
from uuid import uuid4

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from src.service import team_export_details as detail, operations_service
from src.models.admin_ops_requests import ExportRequest
from src.repository.admin_management_repository import AdminManagementRepository
from src.service import export_service
from unittest.mock import AsyncMock

RID=uuid4()
TEAM={'registration_id':RID,'team_name':'Team Example','status':'DRAFT',
      'registration_code':'MIIH-ONE','team_id':None,'members':4,'edition_code':'MIIH_UPCOMING'}
MEMBERS=[{'registration_id':RID,'member_order':i,'full_name':f'Person {i}',
          'email':f'person{i}@example.com','mobile':f'90000000{i:02d}',
          'role':'LEADER' if i==1 else 'MEMBER','photo_filename':f'{i}.png'} for i in range(1,5)]


def test_grouped_excel_cells_expand_and_keep_team_row():
    teams=detail.assemble_teams([TEAM],MEMBERS)
    assert teams[0]['all_member_names']=='Person 1\nPerson 2\nPerson 3\nPerson 4'
    raw=detail.render_team_workbook(teams,MEMBERS,[],[],[],detail.TEAM_COLUMNS,
                                    lambda x: '' if x is None else str(x),super_admin=True)
    ws=load_workbook(BytesIO(raw))['Teams']
    header=[c.value for c in ws[1]]
    idx=header.index('All Member Emails')+1
    cell=ws.cell(2,idx)
    assert cell.value.split('\n') == [f'person{i}@example.com' for i in range(1,5)]
    assert cell.alignment.wrap_text is True
    assert ws.row_dimensions[2].height >= 60
    assert ws.column_dimensions[cell.column_letter].width >= 20
    assert ws.max_row==2


def test_json_has_all_four_members_and_keeps_draft_unsubmitted():
    teams=detail.assemble_teams([TEAM],MEMBERS)
    result=json.loads(detail.render_team_json(teams,MEMBERS,[],[],[],super_admin=True))
    assert len(result['registrations'])==1
    assert len(result['registrations'][0]['members'])==4
    assert result['registrations'][0]['team']['status']=='DRAFT'
    assert result['registrations'][0]['team']['team_id'] is None
    assert len(result['registrations'][0]['payment_evidence'])==0
    regular=json.loads(detail.render_team_json(teams,MEMBERS,[],[],[],super_admin=False))
    assert 'payment_evidence' not in regular['registrations'][0]


def test_json_request_accepted_only_for_teams():
    assert ExportRequest(dataset='TEAMS',file_format='JSON').file_format=='JSON'


class Result:
    def __init__(self, rows):self.rows=rows
    def mappings(self):return self
    def one_or_none(self):return self.rows[0] if self.rows else None
    def scalars(self):return self
    def all(self):return self.rows

class Session:
    def __init__(self):self.queries=[];self.commits=0
    async def execute(self, stmt,params=None):
        sql=str(stmt);self.queries.append((sql,params))
        if 'SELECT id,role,approval_status FROM public.admin_users' in sql:
            return Result([{'id':TARGET,'role':'ADMIN','approval_status':'APPROVED'}])
        if 'SELECT permission_key FROM public.admin_permission_catalog' in sql:
            return Result(['TEAMS_VIEW','DASHBOARD_VIEW'])
        if 'SELECT permission_key FROM public.admin_permission_assignments' in sql:
            return Result(['TEAMS_VIEW'])
        return Result([])
    async def commit(self):self.commits+=1

ACTOR=uuid4();TARGET=uuid4()
@pytest.mark.asyncio
async def test_bulk_permissions_one_commit_and_one_audit():
    session=Session()
    result=await AdminManagementRepository(session).save_access(ACTOR,TARGET,['TEAMS_VIEW'],'ADMIN')
    assert result=='OK' and session.commits==1
    assert len([x for x,_ in session.queries if 'INSERT INTO public.admin_audit_logs' in x])==1
    assert not any('UPDATE public.admin_users SET role' in x for x,_ in session.queries)

@pytest.mark.asyncio
async def test_promoting_super_admin_requires_all_permissions():
    session=Session()
    result=await AdminManagementRepository(session).save_access(ACTOR,TARGET,['TEAMS_VIEW'],'SUPER_ADMIN')
    assert result=='REQUIRES_ALL' and session.commits==0
    session=Session()
    result=await AdminManagementRepository(session).save_access(ACTOR,TARGET,['TEAMS_VIEW','DASHBOARD_VIEW'],'SUPER_ADMIN')
    assert result=='OK' and session.commits==1
    assert any('UPDATE public.admin_users SET role' in x for x,_ in session.queries)

@pytest.mark.asyncio
async def test_revenue_query_excludes_unapproved_and_unverified(monkeypatch):
    captured=[]
    async def fake_many(db,sql,params):captured.append(sql);return []
    monkeypatch.setattr(operations_service,'many',fake_many)
    await operations_service.revenue(None,{'role':'SUPER_ADMIN'})
    assert "r.status='APPROVED'" in captured[0]
    assert "p.status='VERIFIED'" in captured[0]
    assert 'MIIH_TEST_%' in captured[0]

@pytest.mark.asyncio
async def test_draft_payment_cannot_be_verified(monkeypatch):
    async def fake_one(db,sql,params):return {'id':uuid4(),'registration_id':RID,'status':'PENDING'}
    async def fake_scalar(db,sql,params):return 'DRAFT'
    monkeypatch.setattr(operations_service,'one',fake_one)
    monkeypatch.setattr(operations_service,'scalar',fake_scalar)
    class Payload:status='VERIFIED';rejection_reason=None
    with pytest.raises(HTTPException) as ex:
        await operations_service.review_payment(None,{'role':'SUPER_ADMIN'},uuid4(),Payload())
    assert ex.value.status_code==409
